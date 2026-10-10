"""Independent arithmetic oracle: no import from interfaces.py or simulation engine.

Literal expected amounts are separately chosen/hand-checked miniature fixtures.
LibreOffice can recalculate formulas; the source workbook keeps the formulas.
"""
from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DEST = Path(__file__).resolve().parent / "oracle_day1_independent.xlsx"
w = Workbook()
s = w.active
s.title = "Accounting oracle"
s.append(["INDEPENDENT DAY 1 ORACLE — illustrative cases, NOT 104-week results"])
s.append(["Prices A=12 USD, B=12.6 USD; uplift 0.14; fee 25000 once on D2 when locked. Expected column is a separately entered literal."])
s.append(["Scenario", "Option", "A units", "B units", "A price", "B price", "A uplift", "Holding USD", "Lost-sales USD", "Exit flag", "Base purchase USD", "Premium USD", "Exit fee USD", "Formula TCO USD", "Expected USD (literal)", "Difference USD", "Notes"])
cases = [
    ("tiny-base", "D0", 250, 0, 12, 12.6, 0.14, 50, 30, 0, 3500, "D0=100% A; 3000+420+50+30"),
    ("tiny-base", "D1", 150, 100, 12, 12.6, 0.14, 20, 10, 0, 3342, "D1=60/40; 3060+252+20+10"),
    ("tiny-base", "D2", 0, 250, 12, 12.6, 0.14, 10, 0, 1, 28160, "D2 exit; 3150+10+25000; no A premium"),
    ("demand-drop", "D1", 15600, 10400, 12, 12.6, 0.14, 48000, 26552, 0, 419000, "26k purchased vs 22100 realised: 3900 surplus; A floor 15600"),
    ("demand-drop", "D0", 24500, 0, 12, 12.6, 0.14, 76000, 26840, 0, 438000, "Mock D0 purchase quantity is a supplied aggregate, NOT derived from weekly policy"),
]
for row_number, (scenario, option, a, b, pa, pb, uplift, holding, loss, exit_flag, expected, note) in enumerate(cases, start=4):
    s.append([scenario, option, a, b, pa, pb, uplift, holding, loss, exit_flag,
              f"=ROUND(C{row_number}*E{row_number}+D{row_number}*F{row_number},0)",
              f"=ROUND(C{row_number}*E{row_number}*G{row_number},0)",
              f"=IF(J{row_number}=1,25000,0)",
              f"=SUM(H{row_number}:I{row_number},K{row_number}:M{row_number})",
              expected, f"=N{row_number}-O{row_number}", note])
s.append([])
s.append(["EXTRA CHECK", "D1 demand-drop total units", "=C7+D7", "expected", 26000, "surplus", "=C7+D7-22100", "expected", 3900])
s.append(["EXTRA CHECK", "A floor", "=26000*0.6", "expected", 15600, "above rolling floor", "=C7-22100*0.6", "expected", 2340])
s.append(["Boundary", "The 25,936 historical units do NOT independently imply a 26,000 locked forecast. Holding and loss inputs here are stipulated fixture values."])

q = w.create_sheet("Selection oracle")
q.append(["Independent threshold examples; probabilities are fractions, TCO and cash are USD."])
q.append(["Scenario", "Option", "Feasible", "Stockout prob", "P90 cash USD", "TCO USD", "Risk limit", "Cash ceiling", "Eligible?", "Expected scenario outcome (literal)"])
selection = [
    ("tiny-base", "D0", True, 0.09, 3600, 3500, 0.12, 4000, "D1"),
    ("tiny-base", "D1", True, 0.05, 3500, 3342, 0.12, 4000, "D1"),
    ("tiny-base", "D2", True, 0.03, 29000, 28160, 0.12, 4000, "D1"),
    ("no-feasible", "D0", True, 0.09, 3600, 3500, 0.01, 4000, "no_feasible_option"),
    ("no-feasible", "D1", True, 0.05, 3500, 3342, 0.01, 4000, "no_feasible_option"),
    ("no-feasible", "D2", True, 0.03, 29000, 28160, 0.01, 4000, "no_feasible_option"),
]
for row_number, row in enumerate(selection, start=3):
    q.append(list(row[:8]) + [f"=AND(C{row_number},D{row_number}<=G{row_number},E{row_number}<=H{row_number})", row[8]])
q.append(["Policy", "For each scenario: feasible ∩ risk<=limit ∩ cash<=ceiling; min(TCO, Option ID); if none, no_feasible_option."])

p = w.create_sheet("Provenance")
for row in [
    ("Source", "Causora-Day1-Delivery-v4.1(1).zip → causora/API_CONTRACT.md; lib/contracts.ts; demo_data/causora_day1_mock.json"),
    ("Original spec", "Causora.pdf pp. 7–10, 20–22; weekly model and Day 1/Day 2 boundary"),
    ("Independent", "This workbook is composed of human-set miniature cases, fixed literal expected results, and its own cell formulas. It does not import interfaces.py, simulation outputs or mock JSON to calculate expected."),
    ("Trace", "purchase = A units*12 + B units*12.6; premium = A units*12*0.14; exit fee = 25000 only if D2 exits a locked renewal; TCO = sum five lines."),
    ("Limit", "No Monte Carlo stockout/P90 oracle is asserted; illustrative risk/cash literals are only to check selection. No 104-week golden numbers are claimed."),
    ("Rounding", "ROUND to whole USD per cost line; for these cases every premium and purchase is already whole USD, so tie-breaking does not matter."),
    ("Future", "Check date/notice, time horizon vs renewal date, signed base prices, inventory policy, handling procurement floor and end-of-horizon arrivals with the team before extending oracle."),
]:
    p.append(row)
for sheet in w:
    sheet.freeze_panes = "C4" if sheet == s else "C3"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 27
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    sheet["A1"].fill = PatternFill("solid", fgColor="1D3652")
    for cell in sheet[3] if sheet == s else sheet[2] if sheet == q else []:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="286586")
    for col in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 27 if col < sheet.max_column else 65
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="center")
s.column_dimensions["A"].width = 30
s.column_dimensions["Q"].width = 77
q.column_dimensions["J"].width = 35
p.column_dimensions["A"].width = 20
p.column_dimensions["B"].width = 126
w.calculation.fullCalcOnLoad = True
w.calculation.forceFullCalc = True
w.save(DEST)
print(DEST)
